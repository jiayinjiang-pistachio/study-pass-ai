"""RAG 知识库入库：解析 → 清洗 → 分类 → 分块 → 向量化 → 写 Qdrant。

对外只暴露 ingest_file；其余函数均为单文件流水线内部的私有步骤。
"""

import hashlib
import logging
import re
import shutil
import subprocess
import uuid
from pathlib import Path

import pypdfium2 as pdfium
from dashscope import TextEmbedding
from docx import Document
from docx.text.paragraph import Paragraph
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient
from qdrant_client.models import (
    Distance,
    FieldCondition,
    Filter,
    MatchValue,
    PointStruct,
    VectorParams,
)

from ..core import config
from ..core.exceptions import BusinessException, SystemException

logger = logging.getLogger(__name__)

__all__ = ["ingest_file"]

COLLECTION_NAME = "knowledge_chunks"
VECTOR_SIZE = 1024
EMBED_MODEL = "text-embedding-v4"
EMBED_BATCH_SIZE = 10

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50

# 递归切片分隔符，按优先级从结构到字符逐级回退：
# 标题（由 _parse_docx 注入）→ 段落 → 中文句读 → 空格 → 字符兜底。
CHUNK_SEPARATORS = (
    "\n### ",
    "\n## ",
    "\n# ",
    "\n\n",
    "\n",
    "。",
    "！",
    "？",
    "；",
    "，",
    " ",
    "",
)

# 逐页扫描件判定：仅当页面被图片大面积覆盖、且文本层字符很少时才转 OCR。
# 双信号可避免两类误判：纯文本页面被平均字符数拉低（误 OCR），
# 以及扫描页混在正常 PDF 中被整体判定为文本（漏 OCR）。
SCAN_IMAGE_COVERAGE = 0.5
SCAN_TEXT_MAX_CHARS = 200

CATEGORIES = ("resume", "study_material", "general")

# 分类关键字：命中次数计分，显式传入的分类优先于自动判断
_RESUME_KEYWORDS = (
    "简历", "个人简历", "求职意向", "工作经历", "工作经验", "项目经历", "项目经验",
    "教育经历", "教育背景", "实习经历", "技能特长", "专业技能", "自我评价", "联系方式",
    "毕业院校", "主修课程", "工作职责", "荣誉奖项",
    "resume", "work experience", "education", "skills", "projects", "objective",
)
_STUDY_KEYWORDS = (
    "知识点", "课程", "章节", "习题", "例题", "考点", "考试", "复习", "笔记", "课堂",
    "教材", "作业", "题库", "学习目标", "重点难点", "解析", "面试题", "练习题",
    "lecture", "chapter", "exercise", "quiz", "exam", "homework", "study notes",
)

# 零宽字符：PDF / Word 复制粘贴常见，不清掉会干扰分块与向量化
_ZERO_WIDTH = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u2060\ufeff"))
_PAGE_NUMBER_RE = re.compile(r"^\s*(?:[-—–]\s*)?\d{1,5}(?:\s*[-—–])?\s*$")
_SPACE_RE = re.compile(r"[ \t\f\v\u00a0]+")
_BLANK_LINES_RE = re.compile(r"\n{3,}")
# Word 内置标题样式名随语言变化（Heading 1 / 标题 1），style_id 则形如 Heading1
_HEADING_STYLE_RE = re.compile(r"^(?:heading|标题)\s*([1-3])$", re.IGNORECASE)


def _text_chars(text: str) -> int:
    """文本层有效字符数：去掉零宽字符、独立页码行与所有空白后计数。"""
    text = text.translate(_ZERO_WIDTH)
    lines = [
        line
        for line in text.splitlines()
        if not _PAGE_NUMBER_RE.match(line.strip())
    ]
    return len(re.sub(r"\s+", "", "\n".join(lines)))


def _image_coverage(page: pdfium.PdfPage) -> float:
    """图片对象覆盖面积 / 页面面积，钳制到 0..1。

    扫描件的整页位图通常接近 1；正文页的小图标 / 分隔线占比很低。
    """
    width, height = page.get_size()
    area = width * height
    if area <= 0:
        return 0.0

    covered = 0.0
    for obj in page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_IMAGE]):
        try:
            left, bottom, right, top = obj.get_bounds()
            covered += max(0.0, right - left) * max(0.0, top - bottom)
        except Exception:  # 个别畸形对象读不到边界，跳过不影响其余页
            logger.debug("读取图片对象边界失败", exc_info=True)
        finally:
            obj.close()
    return min(covered / area, 1.0)


def _needs_ocr(text: str, image_ratio: float) -> bool:
    """图片覆盖率高且几乎没文字 → 扫描页；其余保留文本层。"""
    return (
        image_ratio >= SCAN_IMAGE_COVERAGE
        and _text_chars(text) < SCAN_TEXT_MAX_CHARS
    )


def _analyze_pdf(path: Path) -> tuple[list[str], list[int]]:
    """逐页取文本层并判定是否需要 OCR，返回 (每页文本, 待 OCR 页下标)。"""
    document = pdfium.PdfDocument(str(path))
    try:
        pages: list[str] = []
        ocr_pages: list[int] = []
        for index in range(len(document)):
            page = document[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                finally:
                    textpage.close()
                need_ocr = _needs_ocr(text, _image_coverage(page))
            finally:
                page.close()
            pages.append(text)
            if need_ocr:
                ocr_pages.append(index)
        return pages, ocr_pages
    finally:
        document.close()


def _parse_pdf(path: Path) -> str:
    """逐页判定：文本层够用时直接取；扫描页单独 OCR 后按页序合并。"""
    pages, ocr_pages = _analyze_pdf(path)
    if not ocr_pages:
        return "\n".join(pages)

    logger.info(
        "PDF 含扫描页，转 OCR：path=%s pages=%s ocr_pages=%s",
        path,
        len(pages),
        len(ocr_pages),
    )
    # 延迟导入：只有扫描页才需要加载 volcengine / pillow，避免普通 PDF 付出导入成本
    from .ocr import ocr_pdf

    recognized = ocr_pdf(path, page_indices=ocr_pages)
    # dict.get(key, default) 的语义：key 存在就取对应值，不存在就返回 default。
    merged = [recognized.get(index, text) for index, text in enumerate(pages)]
    return "\n".join(merged)


def _heading_level(paragraph: Paragraph) -> int:
    """识别 Heading 1 / 2 / 3（兼容「标题 N」样式名）；非标题返回 0。"""
    style = paragraph.style
    name = (style.name if style is not None else "") or ""
    match = _HEADING_STYLE_RE.match(name.strip())
    return int(match.group(1)) if match else 0


def _parse_docx(path: Path) -> str:
    """段落 + 表格单元格；Heading 1/2/3 注入 Markdown 标记供递归切片识别。"""
    document = Document(str(path))
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        level = _heading_level(paragraph)
        # 只给有内容的标题加标记；空段落保留，维持原有段落间距
        if level and text:
            lines.append(f"{'#' * level} {text}")
        else:
            lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _parse_txt(path: Path) -> str:
    """先按 UTF-8（含 BOM）解码，失败回退 GB18030 兼容 GBK 文档。"""
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise BusinessException(
        "TXT 文本编码无法识别，请使用 UTF-8 或 GBK",
        http_status=422,
        detail=f"path={path}",
    )


# .doc 是 OLE 二进制格式，Python 无标准库解析器；按系统可用工具转文本
_DOC_COMMANDS = (
    ("textutil", ["textutil", "-convert", "txt", "-stdout"]),
    ("antiword", ["antiword"]),
    ("soffice", ["soffice", "--headless", "--cat"]),
    ("libreoffice", ["libreoffice", "--headless", "--cat"]),
)


def _parse_doc(path: Path) -> str:
    """调用系统转换工具读取 .doc（macOS textutil / antiword / LibreOffice）。"""
    for binary, command in _DOC_COMMANDS:
        if shutil.which(binary) is None:
            continue
        try:
            result = subprocess.run(
                [*command, str(path)], capture_output=True, timeout=60
            )
        except (OSError, subprocess.TimeoutExpired):
            continue
        if result.returncode != 0:
            continue
        text = ""
        for encoding in ("utf-8", "gb18030"):
            try:
                text = result.stdout.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text.strip():
            return text

    raise SystemException(
        ".doc 解析失败",
        detail="未找到可用转换工具（textutil / antiword / LibreOffice）或文件无法转换",
    )


# 字典分派解析器：扩展名 → 解析函数
_PARSERS = {
    ".pdf": _parse_pdf,
    ".doc": _parse_doc,
    ".docx": _parse_docx,
    ".txt": _parse_txt,
}


def _clean(text: str) -> str:
    """去零宽字符、合并空白、删独立页码行、压缩连续空行。"""
    text = text.translate(_ZERO_WIDTH).replace("\r\n", "\n").replace("\r", "\n")
    lines: list[str] = []
    for raw in text.split("\n"):
        line = _SPACE_RE.sub(" ", raw).strip()
        if _PAGE_NUMBER_RE.match(line):
            continue
        lines.append(line)
    return _BLANK_LINES_RE.sub("\n\n", "\n".join(lines)).strip()


def _classify(text: str) -> str:
    """按关键字命中次数打分；无命中或平局归为 general。"""
    lowered = text.lower()
    resume_score = sum(lowered.count(keyword) for keyword in _RESUME_KEYWORDS)
    study_score = sum(lowered.count(keyword) for keyword in _STUDY_KEYWORDS)
    if resume_score > study_score:
        return "resume"
    if study_score > resume_score:
        return "study_material"
    return "general"


def _chunk_text(text: str) -> list[str]:
    """递归切片：标题 → 段落 → 中文句读 → 字符，尽量在语义边界断开。"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=list(CHUNK_SEPARATORS),
        keep_separator=True,
    )
    return [chunk.strip() for chunk in splitter.split_text(text) if chunk.strip()]


def _dedupe(chunks: list[str]) -> list[str]:
    """按内容 MD5 去重，保留首次出现的位置。"""
    seen: set[str] = set()
    unique: list[str] = []
    for chunk in chunks:
        digest = hashlib.md5(chunk.encode("utf-8"), usedforsecurity=False).hexdigest()
        if digest in seen:
            continue
        seen.add(digest)
        unique.append(chunk)
    return unique


def _embed(texts: list[str]) -> list[list[float]]:
    """每批 10 条调用 DashScope，并按 text_index 还原成与输入一致的下标顺序。"""
    if not config.DASHSCOPE_API_KEY:
        raise SystemException(
            "向量化服务未配置", detail="DASHSCOPE_API_KEY 未配置"
        )

    vectors: list[list[float] | None] = [None] * len(texts)
    for offset in range(0, len(texts), EMBED_BATCH_SIZE):
        batch = texts[offset : offset + EMBED_BATCH_SIZE]
        try:
            response = TextEmbedding.call(
                model=EMBED_MODEL,
                input=batch,
                dimension=VECTOR_SIZE,
                api_key=config.DASHSCOPE_API_KEY,
            )
        except Exception as exc:
            raise SystemException("文本向量化请求失败", detail=str(exc)) from exc

        if response.status_code != 200:
            raise SystemException(
                "文本向量化失败",
                detail=f"status={response.status_code} message={response.message}",
            )

        for item in (response.output or {}).get("embeddings") or []:
            vectors[offset + item["text_index"]] = item["embedding"]

    if any(vector is None for vector in vectors):
        raise SystemException("文本向量化结果不完整", detail=f"expected={len(texts)}")
    return [vector for vector in vectors if vector is not None]


def _write_qdrant(
    payloads: list[dict],
    vectors: list[list[float]],
    resource_id: int | None = None,
    replace: bool = False,
) -> None:
    """确保 collection 存在；replace 时先删同一资源的旧向量，再批量 upsert。"""
    if replace and resource_id is None:
        raise ValueError("replace=True 时必须提供 resource_id")

    try:
        client = QdrantClient(
            host=config.QDRANT_HOST, port=config.QDRANT_PORT, timeout=30
        )
    except Exception as exc:
        raise SystemException("Qdrant 连接失败", detail=str(exc)) from exc

    try:
        if not client.collection_exists(COLLECTION_NAME):
            client.create_collection(
                collection_name=COLLECTION_NAME,
                vectors_config=VectorParams(size=VECTOR_SIZE, distance=Distance.COSINE),
            )

        # 分类变更重新入库前先清理旧向量，避免同一文件留下两套分类的 chunk
        if replace and resource_id is not None:
            client.delete(
                collection_name=COLLECTION_NAME,
                points_selector=Filter(
                    must=[
                        FieldCondition(
                            key="resource_id",
                            match=MatchValue(value=resource_id),
                        )
                    ]
                ),
                wait=True,
            )

        client.upsert(
            collection_name=COLLECTION_NAME,
            points=[
                PointStruct(id=str(uuid.uuid4()), vector=vector, payload=payload)
                for payload, vector in zip(payloads, vectors)
            ],
            wait=True,
        )
    except Exception as exc:
        raise SystemException("知识库写入失败", detail=str(exc)) from exc
    finally:
        client.close()


def ingest_file(
    file_path: str | Path,
    user_id: int,
    doc_category: str | None = None,
    file_name: str | None = None,
    resource_id: int | None = None,
    replace: bool = False,
) -> dict:
    """把单个文件向量化写入 Qdrant，返回入库摘要。

    doc_category 显式传入时优先使用（只接受 resume / study_material / general）；
    file_name 用于保留上传时的原始文件名（落盘路径可能是 MD5 名）；
    resource_id 写入 payload 用于按资源重灌，replace=True 时先删除该资源旧向量。
    """
    path = Path(file_path)
    if not path.is_file():
        raise BusinessException(
            "文件不存在", http_status=404, detail=f"path={path}"
        )

    if doc_category is not None and doc_category not in CATEGORIES:
        raise BusinessException(
            "不支持的文档分类",
            http_status=422,
            detail=f"doc_category={doc_category}",
        )

    parser = _PARSERS.get(path.suffix.lower())
    if parser is None:
        raise BusinessException(
            "仅支持 PDF / DOC / DOCX / TXT 文件",
            http_status=422,
            detail=f"path={path}",
        )

    try:
        raw_text = parser(path)
    except (BusinessException, SystemException):
        raise
    except Exception as exc:
        raise BusinessException(
            "文件解析失败", http_status=422, detail=f"path={path} error={exc}"
        ) from exc

    text = _clean(raw_text)
    if not text:
        raise BusinessException(
            "文件未解析出有效文本", http_status=422, detail=f"path={path}"
        )

    category = doc_category or _classify(text)
    chunks = _dedupe(_chunk_text(text))
    if not chunks:
        raise BusinessException(
            "文件未解析出有效文本", http_status=422, detail=f"path={path}"
        )

    vectors = _embed(chunks)
    resolved_name = file_name or path.name
    payloads = [
        {
            "text": chunk,
            "user_id": user_id,
            "doc_category": category,
            "file_name": resolved_name,
            "chunk_index": index,
            "resource_id": resource_id,
        }
        for index, chunk in enumerate(chunks)
    ]
    _write_qdrant(payloads, vectors, resource_id=resource_id, replace=replace)

    logger.info(
        "知识库入库完成：file=%s user=%s category=%s chunks=%s",
        resolved_name,
        user_id,
        category,
        len(chunks),
    )
    return {
        "file_name": resolved_name,
        "doc_category": category,
        "chunk_count": len(chunks),
        "collection": COLLECTION_NAME,
    }
